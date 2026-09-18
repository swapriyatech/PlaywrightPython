from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


class TestDataEngine:
    def load_json(self, path: Path) -> dict[str, Any]:
        with path.open(encoding="utf-8") as stream:
            value = json.load(stream)
        if not isinstance(value, dict):
            raise ValueError(f"JSON test data must be an object: {path}")
        return value

    def load_csv(self, path: Path) -> list[dict[str, str]]:
        with path.open(newline="", encoding="utf-8") as stream:
            return list(csv.DictReader(stream))

    def load_excel(self, path: Path, sheet: str) -> list[dict[str, Any]]:
        workbook = load_workbook(path, read_only=True, data_only=True)
        if sheet not in workbook.sheetnames:
            raise ValueError(f"Worksheet {sheet!r} not found in {path}")
        rows = workbook[sheet].iter_rows(values_only=True)
        headers = [str(value) for value in next(rows)]
        return [dict(zip(headers, row, strict=True)) for row in rows]

    def merge(self, *sources: dict[str, Any]) -> dict[str, Any]:
        result: dict[str, Any] = {}
        for source in sources:
            result.update(source)
        return result
