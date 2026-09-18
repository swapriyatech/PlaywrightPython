from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook


class ExcelUtilities:
    @staticmethod
    def sheet_names(path: Path) -> list[str]:
        workbook = load_workbook(path, read_only=True, data_only=True)
        return workbook.sheetnames

    @staticmethod
    def read_rows(path: Path, sheet: str) -> list[dict[str, Any]]:
        workbook = load_workbook(path, read_only=True, data_only=True)
        if sheet not in workbook.sheetnames:
            raise ValueError(f"Worksheet {sheet!r} not found: {path}")
        rows = workbook[sheet].iter_rows(values_only=True)
        headers = [str(value) for value in next(rows)]
        return [dict(zip(headers, row, strict=True)) for row in rows]

    @staticmethod
    def read_cell(path: Path, sheet: str, cell: str) -> Any:
        workbook = load_workbook(path, read_only=True, data_only=True)
        if sheet not in workbook.sheetnames:
            raise ValueError(f"Worksheet {sheet!r} not found: {path}")
        return workbook[sheet][cell].value

    @staticmethod
    def write_rows(path: Path, sheet: str, rows: list[dict[str, Any]]) -> None:
        if not rows:
            raise ValueError("At least one row is required")
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = sheet
        headers = list(rows[0])
        worksheet.append(headers)
        for row in rows:
            worksheet.append([row.get(header) for header in headers])
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(path)

    @staticmethod
    def append_rows(path: Path, sheet: str, rows: list[dict[str, Any]]) -> None:
        if not rows:
            return
        workbook = load_workbook(path)
        if sheet not in workbook.sheetnames:
            raise ValueError(f"Worksheet {sheet!r} not found: {path}")
        worksheet = workbook[sheet]
        headers = [cell.value for cell in worksheet[1]]
        for row in rows:
            worksheet.append([row.get(header) for header in headers])
        workbook.save(path)
